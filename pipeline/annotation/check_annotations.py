"""
Checks returned annotation workbooks and computes inter-annotator agreement.

For each annotator workbook (main sheet 'Ulasan', translated sheet 'Ulasan_terjemahan',
dictionary sheet 'Kamus') it verifies that IDs and texts are unchanged, flags invalid or
inconsistent labels, merges the two review sheets into one label per review, reports
progress, and computes Cohen's kappa on the items both annotators have completed.

Merge rule: Indonesian/Manadonese reviews are labelled in 'Ulasan'. Non-English foreign
reviews are labelled in 'Ulasan_terjemahan'. English reviews take the 'Ulasan_terjemahan'
label if present, otherwise the 'Ulasan' label.

Usage (from the pipeline root):
    python annotation/check_annotations.py annotation/tmp/Anotasi_Annotator_A.xlsx annotation/tmp/Anotasi_Annotator_B.xlsx
Writes <first file's folder>/merged_labels.csv
"""
import os
import sys
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

ANNOT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ANNOT_DIR)
from build_translation_addendum import LANGUAGE  # noqa: E402

POLARITY = {'Positif', 'Negatif', 'Netral-campuran'}
YESNO = {'Ya', 'Tidak'}
LABELS = ['mention', 'service_pol', 'overall_pol', 'manado']


def clean(v):
    return v.strip() if isinstance(v, str) and v.strip() else None


def read_reviews(ws, id_col, text_col, first_label_col):
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[id_col]:
            continue
        labels = [clean(r[first_label_col + k]) for k in range(4)]
        rows[str(r[id_col]).strip()] = {'text': r[text_col], **dict(zip(LABELS, labels))}
    return rows


def validate(name, sheet, rows, problems):
    for rid, d in rows.items():
        for col, allowed in (('mention', YESNO), ('service_pol', POLARITY),
                             ('overall_pol', POLARITY), ('manado', YESNO)):
            if d[col] is not None and d[col] not in allowed:
                problems.append(f'{name} {sheet} {rid}: invalid {col} = {d[col]!r}')
        if d['mention'] == 'Tidak' and d['service_pol']:
            problems.append(f'{name} {sheet} {rid}: service polarity filled but mention = Tidak')
        if d['mention'] == 'Ya' and not d['service_pol'] and d['overall_pol']:
            problems.append(f'{name} {sheet} {rid}: mention = Ya but service polarity empty')


def complete(d):
    return bool(d and d['mention'] and d['overall_pol'] and (d['mention'] == 'Tidak' or d['service_pol']))


def load(path, key, problems):
    name = os.path.basename(path)
    wb = load_workbook(path, data_only=True)
    main = read_reviews(wb['Ulasan'], 0, 1, 2)
    trans_name = next((s for s in wb.sheetnames if s.lower().startswith('ulasan_t')), None)
    trans = read_reviews(wb[trans_name], 0, 2, 3) if trans_name else {}

    # IDs and texts must match what was sent
    if list(main) != list(key.index):
        problems.append(f'{name} Ulasan: review IDs or order differ from the template')
    for rid, d in main.items():
        if rid in key.index and str(d['text']) != str(key.loc[rid, 'original_text']):
            problems.append(f'{name} Ulasan {rid}: review text was changed')
    if set(trans) != set(LANGUAGE):
        problems.append(f'{name} {trans_name}: IDs differ from the 72 translated reviews')
    validate(name, 'Ulasan', main, problems)
    validate(name, trans_name, trans, problems)

    merged = {}
    for rid in key.index:
        lang = LANGUAGE.get(rid)
        if lang is None:
            merged[rid] = main.get(rid)
        elif lang == 'Inggris':
            merged[rid] = trans.get(rid) if complete(trans.get(rid)) else main.get(rid)
        else:
            merged[rid] = trans.get(rid)
    foreign_in_main = [r for r in LANGUAGE if LANGUAGE[r] != 'Inggris' and complete(main.get(r))]

    kamus = {}
    for r in wb['Kamus'].iter_rows(min_row=2, values_only=True):
        if r and r[1]:
            kamus[str(r[1])] = {'correct': clean(r[3]), 'type': clean(r[4]), 'fix': clean(r[5])}
    for w, d in kamus.items():
        if d['correct'] and d['correct'] not in {'Ya', 'Tidak', 'Ragu'}:
            problems.append(f'{name} Kamus {w}: invalid value {d["correct"]!r}')
        if d['correct'] == 'Tidak' and not d['fix']:
            problems.append(f'{name} Kamus {w}: marked wrong but no correction given')
    return merged, kamus, foreign_in_main


def kappa(a, b):
    n = len(a)
    if n == 0:
        return float('nan')
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else float('nan')


def main():
    paths = sys.argv[1:3]
    key = pd.read_csv(os.path.join(ANNOT_DIR, 'gold_sample_key.csv')).set_index('review_id')
    problems = []
    (ma, ka, fa), (mb, kb, fb) = (load(p, key, problems) for p in paths)

    print('PROGRESS')
    for tag, m, k, f in (('A', ma, ka, fa), ('B', mb, kb, fb)):
        done = [r for r in key.index if complete(m[r])]
        indo = [r for r in done if r not in LANGUAGE]
        print(f'  Annotator {tag}: {len(done)}/300 reviews complete '
              f'({len(indo)}/228 Indonesian-Manadonese, {len(done) - len(indo)}/72 foreign); '
              f'dictionary {sum(1 for d in k.values() if d["correct"])}/{len(k)}')
        if f:
            print(f'    note: non-English foreign reviews labelled in the main sheet (translated sheet used instead): {f}')

    print('\nPROBLEMS' if problems else '\nPROBLEMS: none')
    for p in problems:
        print('  ' + p)

    both = [r for r in key.index if complete(ma[r]) and complete(mb[r])]
    print(f'\nAGREEMENT on {len(both)} reviews completed by both')
    for col, label in (('mention', 'Mentions service'), ('overall_pol', 'Overall polarity'),
                       ('manado', 'Contains Manadonese')):
        a = [ma[r][col] for r in both]
        b = [mb[r][col] for r in both]
        print(f'  {label:22} agreement {sum(x == y for x, y in zip(a, b)) / len(both):6.1%}   '
              f'kappa {kappa(a, b):.3f}')
    sv = [r for r in both if ma[r]['mention'] == 'Ya' and mb[r]['mention'] == 'Ya']
    a = [ma[r]['service_pol'] for r in sv]
    b = [mb[r]['service_pol'] for r in sv]
    if sv:
        print(f'  {"Service polarity":22} agreement {sum(x == y for x, y in zip(a, b)) / len(sv):6.1%}   '
              f'kappa {kappa(a, b):.3f}   (n = {len(sv)} where both say service is mentioned)')

    kw = [w for w in ka if ka[w]['correct'] and kb.get(w, {}).get('correct')]
    if kw:
        print(f'\nDICTIONARY on {len(kw)} entries rated by both')
        for col, label in (('correct', 'Translation correct'), ('type', 'Word type')):
            # an entry one annotator left blank is not a category: compare only entries both labelled
            ws = [w for w in kw if ka[w][col] and kb[w][col]]
            a = [ka[w][col] for w in ws]
            b = [kb[w][col] for w in ws]
            print(f'  {label:22} agreement {sum(x == y for x, y in zip(a, b)) / len(ws):6.1%}   '
                  f'kappa {kappa(a, b):.3f}   (n = {len(ws)})')
        for tag, k in (('A', ka), ('B', kb)):
            c = Counter(d['correct'] for d in k.values() if d['correct'])
            t = Counter(d['type'] for d in k.values() if d['type'])
            print(f'  Annotator {tag}: correct {dict(c)} | type {dict(t)}')

    rows = []
    for r in key.index:
        row = {'review_id': r, 'language': LANGUAGE.get(r, 'Indonesia/Manado')}
        for tag, m in (('A', ma), ('B', mb)):
            for col in LABELS:
                row[f'{col}_{tag}'] = (m[r] or {}).get(col)
        rows.append(row)
    out = os.path.join(os.path.dirname(os.path.abspath(paths[0])), 'merged_labels.csv')
    pd.DataFrame(rows).to_csv(out, index=False)
    print(f'\nMerged labels written to {out}')


if __name__ == '__main__':
    main()
