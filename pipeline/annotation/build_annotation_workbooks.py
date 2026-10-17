"""
Builds the human-annotation package for the gold-standard evaluation.

Draws a stratified sample of 300 reviews from the 1,095-review analytical sample
(stratified by star rating, oversampling the rare 2- and 3-star reviews) and writes:
  - Anotasi_Annotator_A.xlsx / Anotasi_Annotator_B.xlsx : identical blind workbooks
    (original review text only; no ratings, no pipeline scores)
  - gold_sample_key.csv : review_id -> metadata, kept by the authors for evaluation

Run from the pipeline root:  python annotation/build_annotation_workbooks.py
"""
import json
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR  = os.path.join(BASE_DIR, 'annotation')
OUTPUT   = os.path.join(BASE_DIR, 'data', 'output')
DICT     = os.path.join(BASE_DIR, 'data', 'input', 'manado_slang_dict.json')

SEED = 42
# Reviews per star rating. 2- and 3-star reviews are taken in full because they are rare.
ALLOCATION = {1: 70, 2: 41, 3: 44, 4: 70, 5: 75}
KEY = ['review_author', 'google_id', 'review_date']

FONT     = Font(name='Arial', size=10)
BOLD     = Font(name='Arial', size=10, bold=True)
TITLE    = Font(name='Arial', size=14, bold=True)
HEAD     = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HEADFILL = PatternFill('solid', fgColor='305496')
INPUT    = PatternFill('solid', fgColor='FFF2CC')
THIN     = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
WRAP     = Alignment(wrap_text=True, vertical='top')


def build_sample():
    cleaned = pd.read_csv(os.path.join(OUTPUT, 'reviews_cleaned.csv'))
    translated = pd.read_csv(os.path.join(OUTPUT, 'reviews_translated.csv'))
    normalized = pd.read_csv(os.path.join(OUTPUT, 'reviews_normalized.csv'))
    analysed = pd.read_csv(os.path.join(OUTPUT, 'reviews_preprocessed.csv'))[KEY]

    df = cleaned.rename(columns={'review_text': 'original_text'})
    df['translated_text'] = translated['review_text']
    df['normalized_text'] = normalized['review_text']
    df['dict_substituted'] = df['translated_text'] != df['normalized_text']
    df = df.merge(analysed, on=KEY, how='inner')  # keep only the 1,095 analysed reviews

    parts = [df[df['review_rating'] == r].sample(n=n, random_state=SEED)
             for r, n in ALLOCATION.items()]
    sample = pd.concat(parts).sample(frac=1, random_state=SEED).reset_index(drop=True)
    sample.insert(0, 'review_id', [f'R{i:03d}' for i in range(1, len(sample) + 1)])
    return sample


def header(ws, cols):
    for c, (name, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEADFILL, WRAP, THIN
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = 'A2'


def dropdown(ws, col_letter, options, n_rows):
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    dv.error, dv.errorTitle = 'Pilih salah satu opsi dari daftar.', 'Isian tidak valid'
    ws.add_data_validation(dv)
    dv.add(f'{col_letter}2:{col_letter}{n_rows + 1}')


def guide_sheet(wb, n_reviews, n_dict):
    ws = wb.active
    ws.title = 'Panduan'
    ws.column_dimensions['A'].width = 34
    for col in 'BCDE':
        ws.column_dimensions[col].width = 22
    rows = [
        ('Panduan Anotasi Ulasan Restoran', TITLE),
        ('', None),
        ('Terima kasih sudah membantu! Isi DUA sheet: "Ulasan" dan "Kamus". Hanya isi sel berwarna kuning.', FONT),
        ('Kerjakan SENDIRI. Jangan berdiskusi dengan annotator lain dan jangan mencari bantuan di Google/AI.', BOLD),
        ('Nilai berdasarkan isi teks saja. Jika ragu, pilih jawaban yang paling masuk akal dan tulis di kolom Catatan.', FONT),
        ('', None),
        ('SHEET "Ulasan": 4 kolom yang diisi', BOLD),
        ('1. Bahas pelayanan?', BOLD),
        ('   Ya = ulasan menyebut pelayanan/staf: keramahan, kecepatan, lama menunggu, pemesanan, pembayaran, penyajian, perhatian pelayan.', FONT),
        ('   Tidak = hanya membahas makanan, rasa, harga, tempat, suasana, kebersihan, parkir, dll.', FONT),
        ('2. Sentimen pelayanan (HANYA jika kolom 1 = Ya; jika Tidak, biarkan kosong)', BOLD),
        ('   Positif / Negatif / Netral-campuran: penilaian penulis terhadap PELAYANAN saja.', FONT),
        ('3. Sentimen keseluruhan', BOLD),
        ('   Positif / Negatif / Netral-campuran: kesan keseluruhan penulis terhadap restoran.', FONT),
        ('   Netral-campuran = tidak ada penilaian jelas, atau pujian dan keluhan seimbang.', FONT),
        ('4. Ada bahasa Manado?', BOLD),
        ('   Ya = ada minimal satu kata khas Manado (mis. nyanda, torang, sadap, so, mar, kong, dang, ngana).', FONT),
        ('   Slang umum Indonesia (gak, bgt, yg, aja) TIDAK dihitung sebagai bahasa Manado.', FONT),
        ('', None),
        ('Contoh pengisian:', BOLD),
    ]
    for r, (text, font) in enumerate(rows, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        if font:
            cell.font = font

    ex_head = ['Ulasan', 'Bahas pelayanan?', 'Sentimen pelayanan', 'Sentimen keseluruhan', 'Ada bahasa Manado?']
    examples = [
        ['Pelayan ramah skali, makanan lumayan.', 'Ya', 'Positif', 'Positif', 'Ya'],
        ['Tempatnya nyaman, cocok for kumpul keluarga.', 'Tidak', '', 'Positif', 'Ya'],
        ['Makanan sadap, mar pelayanan lambat noh, 30 menit baru datang.', 'Ya', 'Negatif', 'Netral-campuran', 'Ya'],
        ['Harganya mahal bgt, rasanya biasa aja.', 'Tidak', '', 'Negatif', 'Tidak'],
    ]
    start = len(rows) + 1
    for c, name in enumerate(ex_head, start=1):
        cell = ws.cell(row=start, column=c, value=name)
        cell.font, cell.fill, cell.border, cell.alignment = HEAD, HEADFILL, THIN, WRAP
    for i, ex in enumerate(examples, start=start + 1):
        for c, v in enumerate(ex, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.border, cell.alignment = FONT, THIN, WRAP

    r = start + len(examples) + 2
    more = [
        ('SHEET "Kamus": cek terjemahan kata', BOLD),
        ('   Terjemahan benar? Ya / Tidak / Ragu: apakah padanan Bahasa Indonesianya tepat dalam konteks ulasan restoran.', FONT),
        ('   Jika Tidak: tulis terjemahan yang benar di kolom Koreksi.', FONT),
        ('   Jenis kata: Bahasa Manado / Slang-singkatan umum / Bukan slang (kata baku/lainnya).', FONT),
        ('   Contoh: nyanda → tidak = Ya, Bahasa Manado.  gak → tidak = Ya, Slang-singkatan umum.', FONT),
        ('', None),
        ('Progres', BOLD),
    ]
    for text, font in more:
        cell = ws.cell(row=r, column=1, value=text)
        if font:
            cell.font = font
        r += 1
    ws.cell(row=r, column=1, value='Ulasan selesai (kolom Sentimen keseluruhan terisi)').font = FONT
    ws.cell(row=r, column=2, value=f'=COUNTA(Ulasan!E2:E{n_reviews + 1})&" / {n_reviews}"').font = BOLD
    ws.cell(row=r + 1, column=1, value='Kata kamus selesai').font = FONT
    ws.cell(row=r + 1, column=2, value=f'=COUNTA(Kamus!D2:D{n_dict + 1})&" / {n_dict}"').font = BOLD


def review_sheet(wb, sample):
    ws = wb.create_sheet('Ulasan')
    header(ws, [('ID', 7), ('Ulasan', 70), ('Bahas pelayanan?', 13), ('Sentimen pelayanan', 16),
                ('Sentimen keseluruhan', 16), ('Ada bahasa Manado?', 13), ('Catatan (opsional)', 28)])
    for i, row in enumerate(sample.itertuples(), start=2):
        ws.cell(row=i, column=1, value=row.review_id).font = FONT
        ws.cell(row=i, column=2, value=row.original_text).font = FONT
        for c in range(1, 8):
            cell = ws.cell(row=i, column=c)
            cell.alignment, cell.border = WRAP, THIN
            cell.font = FONT
            if c >= 3:
                cell.fill = INPUT
    n = len(sample)
    dropdown(ws, 'C', ['Ya', 'Tidak'], n)
    dropdown(ws, 'D', ['Positif', 'Negatif', 'Netral-campuran'], n)
    dropdown(ws, 'E', ['Positif', 'Negatif', 'Netral-campuran'], n)
    dropdown(ws, 'F', ['Ya', 'Tidak'], n)


def dict_sheet(wb, entries):
    ws = wb.create_sheet('Kamus')
    header(ws, [('No', 6), ('Kata', 16), ('Diterjemahkan jadi', 20), ('Terjemahan benar?', 15),
                ('Jenis kata', 24), ('Koreksi (jika salah)', 22), ('Catatan (opsional)', 28)])
    for i, (src, tgt) in enumerate(entries, start=2):
        for c, v in enumerate([i - 1, src, tgt], start=1):
            ws.cell(row=i, column=c, value=v)
        for c in range(1, 8):
            cell = ws.cell(row=i, column=c)
            cell.font, cell.border, cell.alignment = FONT, THIN, WRAP
            if c >= 4:
                cell.fill = INPUT
    n = len(entries)
    dropdown(ws, 'D', ['Ya', 'Tidak', 'Ragu'], n)
    dropdown(ws, 'E', ['Bahasa Manado', 'Slang-singkatan umum', 'Bukan slang'], n)


def main():
    sample = build_sample()
    with open(DICT, encoding='utf-8') as f:
        entries = list(json.load(f).items())

    for tag in ('A', 'B'):
        wb = Workbook()
        guide_sheet(wb, len(sample), len(entries))
        review_sheet(wb, sample)
        dict_sheet(wb, entries)
        wb.save(os.path.join(OUT_DIR, f'Anotasi_Annotator_{tag}.xlsx'))

    key_cols = ['review_id', 'place_name', 'google_id', 'review_author', 'review_date',
                'review_rating', 'dict_substituted', 'original_text', 'translated_text', 'normalized_text']
    sample[key_cols].to_csv(os.path.join(OUT_DIR, 'gold_sample_key.csv'), index=False)
    print(f'{len(sample)} reviews, {len(entries)} dictionary entries')
    print(sample['review_rating'].value_counts().sort_index().to_dict(),
          '| substituted:', int(sample['dict_substituted'].sum()))


if __name__ == '__main__':
    main()
