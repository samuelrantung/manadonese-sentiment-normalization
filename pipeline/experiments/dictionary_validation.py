"""
Native-speaker validation of the normalization dictionary, and what the corrections would change.

Reads the 'Kamus' sheet of both annotation workbooks and reports
  1. agreement on translation correctness and word type (raw agreement and Cohen's kappa)
  2. the entry types both annotators agree on, and the share of the dictionary's substitutions
     that comes from each type
  3. the effect of applying the annotators' corrections (the twelve possessive '-nya' entries
     removed, mantap -> mantap, sadap -> sedap, dpe -> punya dia): reviews whose M or X changes,
     gold-set accuracy on the agreed labels, and the correlation of M with the star rating
The lexicon is frozen: the corrected dictionary is a sensitivity check, not a new condition.

Report: data/output/dictionary_validation.txt
Run from the pipeline root:  python experiments/dictionary_validation.py
"""
import contextlib
import io
import os
import re
import sys
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ANNOT_DIR = os.path.join(BASE_DIR, 'annotation')
sys.path.insert(0, os.path.join(BASE_DIR, 'steps'))
sys.path.insert(0, os.path.join(BASE_DIR, 'experiments'))
sys.path.insert(0, ANNOT_DIR)

from ablation import INPUT, KEY, OUTPUT, PROV, Preprocessor, load_json, normalize   # noqa: E402
from check_annotations import clean, kappa                                     # noqa: E402
from step7_sentiment_scoring import (                                          # noqa: E402
    NEG_LEXICON_FILE, POS_LEXICON_FILE, SERVICE_KEYWORDS_FILE, calculate_sentiment_score,
    classify_sentiment, load_lexicon, load_service_keywords, score_absa_service_quality)

REPORT = os.path.join(OUTPUT, 'dictionary_validation.txt')
TYPES = {'Bahasa Manado': 'Manadonese', 'Slang-singkatan umum': 'general slang', 'Bukan slang': 'not slang'}
POL = {'Positif': 1, 'Netral-campuran': 0, 'Negatif': -1}


def read_kamus(path):
    ws = load_workbook(path, read_only=True, data_only=True)['Kamus']
    return {str(r[1]).strip().lower(): {'correct': clean(r[3]), 'type': clean(r[4])}
            for r in ws.iter_rows(min_row=2, values_only=True) if r and r[1]}


def corrected(proposed):
    lex = {k: v for k, v in proposed.items() if not k.endswith('nya')}
    lex.update({'mantap': 'mantap', 'mantapp': 'mantap', 'sadap': 'sedap', 'sadapp': 'sedap',
                'dpe': 'punya dia', 'depe': 'punya dia'})
    return lex


def main():
    out = []
    p = out.append
    ka = read_kamus(os.path.join(ANNOT_DIR, 'Anotasi_Annotator_A.xlsx'))
    kb = read_kamus(os.path.join(ANNOT_DIR, 'Anotasi_Annotator_B.xlsx'))
    proposed = load_json(os.path.join(INPUT, 'manado_slang_dict.json'))
    words = sorted(proposed)

    # ---- 1. agreement ----
    p(f'DICTIONARY VALIDATION ({len(words)} entries, two annotators)\n')
    p('1. AGREEMENT')
    a = [ka[w]['correct'] for w in words]
    b = [kb[w]['correct'] for w in words]
    both_ok = sum(x == 'Ya' and y == 'Ya' for x, y in zip(a, b))
    p(f'  Translation correct: A {dict(Counter(a))} | B {dict(Counter(b))}')
    p(f'  Accepted by both: {both_ok} of {len(words)} ({both_ok / len(words):.1%}); '
      f'raw agreement {sum(x == y for x, y in zip(a, b)) / len(words):.1%}, kappa {kappa(a, b):.3f}')
    typed = [w for w in words if ka[w]['type'] and kb[w]['type']]
    ta = [ka[w]['type'] for w in typed]
    tb = [kb[w]['type'] for w in typed]
    p(f'  Word type (n = {len(typed)} typed by both): raw agreement '
      f'{sum(x == y for x, y in zip(ta, tb)) / len(typed):.1%}, kappa {kappa(ta, tb):.3f}')
    p(f'  Entries without a type from one annotator: {[w for w in words if w not in typed]}')

    # ---- 2. agreed types and substitutions by type ----
    def agreed_type(w):
        x, y = ka[w]['type'], kb[w]['type']
        return TYPES.get(x, x) if x and x == y else 'disputed or unlabeled'
    etype = {w: agreed_type(w) for w in words}
    p('\n2. ENTRY TYPES (agreed by both annotators)')
    for t, n in Counter(etype.values()).most_common():
        p(f'  {t:22} {n:3}')

    translated = pd.read_csv(os.path.join(OUTPUT, 'reviews_translated.csv'))
    analysed = pd.read_csv(os.path.join(OUTPUT, 'reviews_preprocessed.csv'))
    df = translated.merge(analysed[KEY], on=KEY, how='inner').reset_index(drop=True)
    pattern = re.compile(r'\b(' + '|'.join(map(re.escape, proposed)) + r')\b', re.IGNORECASE)
    subs = Counter()
    for t in df['review_text']:
        text = re.sub(r'\bsalah satu\b', 'satu', str(t).lower())
        subs.update(etype[m.lower()] for m in pattern.findall(text))
    total = sum(subs.values())
    p(f'  Substitutions ({total:,}) by entry type:')
    for t, n in subs.most_common():
        p(f'    {t:22} {n:5} ({n / total:.1%})')

    # ---- 3. effect of the corrections ----
    with contextlib.redirect_stderr(io.StringIO()):
        lexicon = load_lexicon(POS_LEXICON_FILE, NEG_LEXICON_FILE)
        keywords = load_service_keywords(SERVICE_KEYWORDS_FILE)
    prep = Preprocessor()

    def score(lex):
        pat = re.compile(r'\b(' + '|'.join(map(re.escape, lex)) + r')\b', re.IGNORECASE)
        texts = [prep(normalize(t, lex, pat)[0]) for t in df['review_text']]
        return pd.DataFrame({
            'm': [classify_sentiment(calculate_sentiment_score(s, lexicon)) for s in texts],
            'x': [score_absa_service_quality(s, keywords, lexicon) for s in texts]})

    base, corr = score(proposed), score(corrected(proposed))
    assert (base['m'].values == analysed['review_text'].map(
        lambda s: classify_sentiment(calculate_sentiment_score(s, lexicon))).values).all(), \
        'proposed condition does not reproduce the main pipeline'
    y = df['review_rating'].astype(float)
    p('\n3. EFFECT OF APPLYING THE CORRECTIONS (all 1,096 reviews)')
    p(f'  Reviews whose overall polarity M changes: {(base.m != corr.m).sum()}')
    p(f'  Reviews whose service polarity X changes: {(base.x != corr.x).sum()}')
    p(f'  r(M, rating): proposed {base.m.corr(y):.4f} -> corrected {corr.m.corr(y):.4f}')
    p(f'  r(X, rating): proposed {base.x.corr(y):.4f} -> corrected {corr.x.corr(y):.4f}')

    labels = pd.read_csv(os.path.join(ANNOT_DIR, 'merged_labels.csv'))
    key = pd.read_csv(os.path.join(ANNOT_DIR, 'gold_sample_key.csv'))[['review_id'] + KEY]
    pred = df[KEY].assign(m_p=base.m, m_c=corr.m, x_p=base.x, x_c=corr.x)
    g = labels.merge(key, on='review_id').merge(pred, on=KEY)
    agree = lambda c: g[f'{c}_A'].where(g[f'{c}_A'] == g[f'{c}_B'])  # noqa: E731
    go, gs = agree('overall_pol').map(POL), agree('service_pol').map(POL)
    o, s = go.notna(), (agree('mention') == 'Ya') & gs.notna()
    p(f'  Gold overall accuracy (n = {o.sum()}): proposed {(g.m_p[o] == go[o]).mean():.3f} '
      f'-> corrected {(g.m_c[o] == go[o]).mean():.3f}')
    p(f'  Gold service accuracy (n = {s.sum()}): proposed {(g.x_p[s] == gs[s]).mean():.3f} '
      f'-> corrected {(g.x_c[s] == gs[s]).mean():.3f}')
    p('  (Removing the -nya entries changes which tokens survive stop-word removal, which runs')
    p('   before stemming: tempatnya -> tempat is removed as a stop-word, tempatnya is kept and')
    p('   stemmed to tempat. Token positions, negation scope and service windows shift with it.)')

    # ---- 4. stability across the three re-runs ----
    p('\n4. RE-RUN STABILITY (entries compared with the original dictionary)')
    runs = {f'run{i}': load_json(os.path.join(PROV, f'run{i}.json')) for i in (1, 2, 3)}
    corpus_words = set(' '.join(df['review_text'].astype(str).str.lower()).split())
    for name, lex in runs.items():
        a, b = set(proposed), set(lex)
        ac, bc = a & corpus_words, b & corpus_words
        shared = a & b
        same = sum(proposed[w] == lex[w] for w in shared)
        p(f'  {name}: {len(lex)} entries; Jaccard {len(a & b) / len(a | b):.2f} '
          f'(corpus words only {len(ac & bc) / len(ac | bc):.2f}); shared {len(shared)}, '
          f'same target in {same / len(shared):.0%}')
    common = set(proposed).intersection(*runs.values())
    p(f'  Entries in all four dictionaries: {len(common)} '
      f'({len(common & corpus_words)} occur in the corpus): {", ".join(sorted(common))}')

    text = '\n'.join(out)
    with open(REPORT, 'w', encoding='utf-8') as f:
        f.write(text + '\n')
    print(text)


if __name__ == '__main__':
    main()
